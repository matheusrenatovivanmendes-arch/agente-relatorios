from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from models import Paciente, formatar_cpf
from extensions import session
from schemas.schemas import BuscarPacienteSchema
from recursos import caminho_recurso
import locale
from datetime import datetime

# Personalize esses dados pra sua clínica. As imagens não vêm no repositório
# (dado pessoal) — coloque sua logo e assinatura em imagem/logo.png e
# imagem/assinatura.png.
NOME_PROFISSIONAL = "Nome do profissional"
DADOS_PROFISSIONAL = "Profissão e registro profissional"
CIDADE = "Sua Cidade"


def buscar_paciente(dados: BuscarPacienteSchema):
    paciente = session.query(Paciente).filter_by(cpf=formatar_cpf(dados.cpf)).first()
    if not paciente:
       raise ValueError('paciente nao existe')
    return paciente

def texto_evolucao(evolucao: str):
    if not isinstance(evolucao, str) or not evolucao.strip():
        return 'precisa falar da evoluçao'
    evolucao_world = evolucao
    return evolucao_world

def texto_frequencia(frequencia:str):
    if not isinstance(frequencia, str) or not frequencia.strip():
        return 'precisa falar da frequencia'
    frequencia_world = frequencia
    return frequencia_world

def texto_consideracoes(consideracoes:str):
    if not isinstance(consideracoes, str) or not consideracoes.strip():
        return 'precisa falar das consideraçoes'
    consideracoes_world = consideracoes
    return consideracoes_world

def criar_word(paciente,evolucao_world,frequencia_world,consideracoes_world):
    locale.setlocale(locale.LC_TIME, "pt_BR.utf-8")
    data_atual = datetime.now()
    data_formatada = data_atual.strftime("%d de %B de %Y")
    documento = Document()

    # Margem superior precisa ser maior que a altura da imagem do header (~0.75in),
    # senão o corpo do texto começa a ser desenhado antes da imagem terminar
    secao = documento.sections[0]
    secao.header_distance = Inches(0.3)
    secao.top_margin = Inches(1.3)
    cabecalho = secao.header

    # Pega o parágrafo padrão do cabeçalho, centraliza e insere a logo diretamente nele
    p_img = cabecalho.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.add_run().add_picture(caminho_recurso('imagem/logo.png'), width=Inches(2.0))


    paragrafo_dados = documento.add_paragraph()
    run = paragrafo_dados.add_run(
        f"DADOS DE IDENTIFICAÇÃO\n"
        f"NOME: {paciente.nome}\n"
        f"IDADE: {paciente.idade} anos"
    )
    run.font.bold = True
    run.font.size = Pt(12)
    run.font.name = 'Arial'

    paragrafo_titulo = documento.add_paragraph()
    run_titulo = paragrafo_titulo.add_run("Relatório Fonoaudiológica")
    run_titulo.font.size = Pt(12)
    run_titulo.font.name = 'Arial'
    run_titulo.font.bold = True
    paragrafo_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paragrafo1 = documento.add_paragraph()
    paragrafo1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    texto = (f'{paciente.diagnostico}'
             )
    run3 = paragrafo1.add_run(texto)
    run3.font.size = Pt(12)
    run3.font.name = 'Arial'

    paragrafo_evoluçao = documento.add_paragraph()
    run_evolucao = paragrafo_evoluçao.add_run("EVOLUÇÃO")
    run_evolucao.font.size = Pt(12)
    run_evolucao.font.name = 'Arial'
    run_evolucao.font.bold = True
    texto_evolucao = (
        f'{evolucao_world}'
    )
    paragrafo_evolucao_texto = documento.add_paragraph()
    run_evolucao_texto = paragrafo_evolucao_texto.add_run(
        texto_evolucao.replace("\n", " ").replace(". ", ".\n")
    )
    run_evolucao_texto.font.size = Pt(12)
    run_evolucao_texto.font.name = 'Arial'

    
    paragrafo_frequencia = documento.add_paragraph()
    run_frequencia = paragrafo_frequencia.add_run("FREQUÊNCIA TERAPÊUTICA")
    run_frequencia.font.size = Pt(12)
    run_frequencia.font.name = 'Arial'
    run_frequencia.font.bold = True
    texto_frequencia = (
        f'{frequencia_world}'
    )
    paragrafo_frequencia_texto = documento.add_paragraph()
    run_frequencia_texto = paragrafo_frequencia_texto.add_run(
        texto_frequencia.replace("\n", " ").replace(". ", ".\n")
    )
    run_frequencia_texto.font.size = Pt(12)
    run_frequencia_texto.font.name = 'Arial'

    #Considerações Finais
    paragrafo_consideracoes = documento.add_paragraph()
    run_consideracoes = paragrafo_consideracoes.add_run("CONSIDERAÇÕES FINAIS")
    run_consideracoes.font.size = Pt(12)
    run_consideracoes.font.name = 'Arial'
    run_consideracoes.font.bold = True
    texto_consideracoes = (
        f'{consideracoes_world}'
    )
    paragrafo_consideracoes_texto = documento.add_paragraph()
    run_consideracoes_texto = paragrafo_consideracoes_texto.add_run(
        texto_consideracoes.replace("\n", " ").replace(". ", ".\n")
    )
    run_consideracoes_texto.font.size = Pt(12)
    run_consideracoes_texto.font.name = 'Arial'

    #Data
    paragrafo_data = documento.add_paragraph()
    run_data = paragrafo_data.add_run(f"{CIDADE}, {data_formatada}")
    run_data.font.size = Pt(12)
    run_data.font.name = 'Arial'
    run_data.font.bold = True
    #Assinatura
    paragrafo_assinatura = documento.add_paragraph()
    paragrafo_assinatura.add_run().add_picture(
        caminho_recurso('imagem/assinatura.png'),
        width=Inches(1.8)
    )

    #Informaçoes fonoaudiologa
    paragrafo_fonoaudiologa = documento.add_paragraph()
    run_fonoaudiologa1 = paragrafo_fonoaudiologa.add_run(f"{NOME_PROFISSIONAL}\n")
    run_fonoaudiologa1.font.size = Pt(12)
    run_fonoaudiologa1.font.name = 'Arial'
    run_fonoaudiologa1.font.bold = True
    run_fonoaudiologa2 = paragrafo_fonoaudiologa.add_run(DADOS_PROFISSIONAL)
    run_fonoaudiologa2.font.size = Pt(12)
    run_fonoaudiologa2.font.name = 'Arial'
    run_fonoaudiologa2.font.bold = True
    #Observações Complementares (texto padrão, igual pra todos os relatórios)
    paragrafo_observacoes = documento.add_paragraph()
    run_observacoes = paragrafo_observacoes.add_run("OBSERVAÇÕES COMPLEMENTARES")
    run_observacoes.font.size = Pt(12)
    run_observacoes.font.name = 'Arial'
    run_observacoes.font.bold = True

    paragrafo_sigilo = documento.add_paragraph()
    run_sigilo = paragrafo_sigilo.add_run(
        "O presente documento é de caráter sigiloso, com finalidade clínica, cujo a "
        "profissional não se responsabiliza pelo uso indevido do mesmo. Caso seja "
        "necessário informações complementares sobre o caso, a mesma coloca-se a "
        "disposição."
    )
    run_sigilo.font.size = Pt(12)
    run_sigilo.font.name = 'Arial'

    buffer = BytesIO()
    documento.save(buffer)
    buffer.seek(0)
    return buffer


